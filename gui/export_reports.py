"""
CyberSafe Report Export Module
Exports network monitoring data to PDF and Word formats
"""

from datetime import datetime
from pathlib import Path
import json

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    HAS_REPORTLAB = True
except:
    HAS_REPORTLAB = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_PYTHON_DOCX = True
except:
    HAS_PYTHON_DOCX = False


class ReportExporter:
    """Export monitoring data to various formats"""
    
    def __init__(self):
        self.export_dir = Path.home() / "CyberSafe_Reports"
        self.export_dir.mkdir(exist_ok=True)
    
    def export_to_pdf(self, data: dict, filename: str = None) -> bool:
        """Export data to PDF format"""
        if not HAS_REPORTLAB:
            return False
        
        try:
            if filename is None:
                filename = f"CyberSafe_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            
            filepath = self.export_dir / filename
            
            # Create PDF
            doc = SimpleDocTemplate(str(filepath), pagesize=letter)
            elements = []
            
            # Styles
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#00d4ff'),
                spaceAfter=30,
                alignment=1  # Center
            )
            
            heading_style = ParagraphStyle(
                'CustomHeading',
                parent=styles['Heading2'],
                fontSize=14,
                textColor=colors.HexColor('#00d4ff'),
                spaceAfter=12,
                spaceBefore=12
            )
            
            # Title
            elements.append(Paragraph("CyberSafe Network Monitoring Report", title_style))
            elements.append(Spacer(1, 0.3*inch))
            
            # Report Info
            info_data = [
                ["Report Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
                ["Report Type", "Network Monitoring Analysis"],
            ]
            
            info_table = Table(info_data, colWidths=[2*inch, 4*inch])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#1a1a1a')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#444444')),
            ]))
            elements.append(info_table)
            elements.append(Spacer(1, 0.3*inch))
            
            # Statistics Section
            if 'statistics' in data:
                elements.append(Paragraph("Network Statistics", heading_style))
                
                stats = data['statistics']
                stats_data = [
                    ["Metric", "Value"],
                    ["Total Packets Captured", str(stats.get('total_packets', 0))],
                    ["Threats Detected", str(stats.get('threats_detected', 0))],
                    ["Active Connections", str(stats.get('active_connections', 0))],
                    ["Average Bandwidth", f"{stats.get('bandwidth', 0):.2f} MB/s"],
                    ["Monitoring Duration", stats.get('duration', 'N/A')],
                ]
                
                stats_table = Table(stats_data, colWidths=[3*inch, 3*inch])
                stats_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#00d4ff')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#444444')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#1a1a1a'), colors.HexColor('#0a0a0a')]),
                ]))
                elements.append(stats_table)
                elements.append(Spacer(1, 0.3*inch))
            
            # Threats Section
            if 'threats' in data and data['threats']:
                elements.append(Paragraph("Detected Threats", heading_style))
                
                threats_data = [["Type", "Count", "Severity"]]
                for threat in data['threats']:
                    threats_data.append([
                        threat.get('type', 'Unknown'),
                        str(threat.get('count', 0)),
                        threat.get('severity', 'Medium')
                    ])
                
                threats_table = Table(threats_data, colWidths=[2*inch, 2*inch, 2*inch])
                threats_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#ff4444')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#444444')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#1a1a1a'), colors.HexColor('#0a0a0a')]),
                ]))
                elements.append(threats_table)
                elements.append(Spacer(1, 0.3*inch))
            
            # Ports Section
            if 'ports' in data and data['ports']:
                elements.append(Paragraph("Active Ports", heading_style))
                
                ports_data = [["Port", "Protocol", "Service", "Status"]]
                for port in data['ports'][:10]:  # Limit to 10
                    ports_data.append([
                        str(port.get('port', 'N/A')),
                        port.get('protocol', 'N/A'),
                        port.get('service', 'Unknown'),
                        port.get('status', 'Open')
                    ])
                
                ports_table = Table(ports_data, colWidths=[1.5*inch, 1.5*inch, 2*inch, 1.5*inch])
                ports_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#44ff44')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#444444')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#1a1a1a'), colors.HexColor('#0a0a0a')]),
                ]))
                elements.append(ports_table)
            
            # Generate PDF
            doc.build(elements)
            return True
            
        except Exception as e:
            print(f"Error generating PDF: {e}")
            return False
    
    def export_to_word(self, data: dict, filename: str = None) -> bool:
        """Export data to Word format"""
        if not HAS_PYTHON_DOCX:
            return False
        
        try:
            if filename is None:
                filename = f"CyberSafe_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            
            filepath = self.export_dir / filename
            
            # Create Document
            doc = Document()
            
            # Title
            title = doc.add_heading('CyberSafe Network Monitoring Report', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_format = title.runs[0]
            title_format.font.color.rgb = RGBColor(0, 212, 255)
            
            # Report Info
            doc.add_paragraph()
            doc.add_paragraph(f"Report Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph(f"Report Type: Network Monitoring Analysis")
            
            # Statistics Section
            if 'statistics' in data:
                doc.add_heading('Network Statistics', 1)
                
                stats = data['statistics']
                table = doc.add_table(rows=6, cols=2)
                table.style = 'Light Grid Accent 1'
                
                table.rows[0].cells[0].text = "Metric"
                table.rows[0].cells[1].text = "Value"
                
                table.rows[1].cells[0].text = "Total Packets Captured"
                table.rows[1].cells[1].text = str(stats.get('total_packets', 0))
                
                table.rows[2].cells[0].text = "Threats Detected"
                table.rows[2].cells[1].text = str(stats.get('threats_detected', 0))
                
                table.rows[3].cells[0].text = "Active Connections"
                table.rows[3].cells[1].text = str(stats.get('active_connections', 0))
                
                table.rows[4].cells[0].text = "Average Bandwidth"
                table.rows[4].cells[1].text = f"{stats.get('bandwidth', 0):.2f} MB/s"
                
                table.rows[5].cells[0].text = "Monitoring Duration"
                table.rows[5].cells[1].text = stats.get('duration', 'N/A')
                
                doc.add_paragraph()
            
            # Threats Section
            if 'threats' in data and data['threats']:
                doc.add_heading('Detected Threats', 1)
                
                table = doc.add_table(rows=len(data['threats']) + 1, cols=3)
                table.style = 'Light Grid Accent 1'
                
                table.rows[0].cells[0].text = "Type"
                table.rows[0].cells[1].text = "Count"
                table.rows[0].cells[2].text = "Severity"
                
                for i, threat in enumerate(data['threats'], 1):
                    table.rows[i].cells[0].text = threat.get('type', 'Unknown')
                    table.rows[i].cells[1].text = str(threat.get('count', 0))
                    table.rows[i].cells[2].text = threat.get('severity', 'Medium')
                
                doc.add_paragraph()
            
            # Ports Section
            if 'ports' in data and data['ports']:
                doc.add_heading('Active Ports', 1)
                
                ports_list = data['ports'][:10]  # Limit to 10
                table = doc.add_table(rows=len(ports_list) + 1, cols=4)
                table.style = 'Light Grid Accent 1'
                
                table.rows[0].cells[0].text = "Port"
                table.rows[0].cells[1].text = "Protocol"
                table.rows[0].cells[2].text = "Service"
                table.rows[0].cells[3].text = "Status"
                
                for i, port in enumerate(ports_list, 1):
                    table.rows[i].cells[0].text = str(port.get('port', 'N/A'))
                    table.rows[i].cells[1].text = port.get('protocol', 'N/A')
                    table.rows[i].cells[2].text = port.get('service', 'Unknown')
                    table.rows[i].cells[3].text = port.get('status', 'Open')
            
            # Save Document
            doc.save(str(filepath))
            return True
            
        except Exception as e:
            print(f"Error generating Word document: {e}")
            return False
    
    def export_to_json(self, data: dict, filename: str = None) -> bool:
        """Export data to JSON format"""
        try:
            if filename is None:
                filename = f"CyberSafe_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
            
            filepath = self.export_dir / filename
            
            report_data = {
                'report_date': datetime.now().isoformat(),
                'data': data
            }
            
            with open(filepath, 'w') as f:
                json.dump(report_data, f, indent=2, default=str)
            
            return True
        except Exception as e:
            print(f"Error generating JSON report: {e}")
            return False
    
    def get_export_directory(self) -> str:
        """Get the reports export directory"""
        return str(self.export_dir)
